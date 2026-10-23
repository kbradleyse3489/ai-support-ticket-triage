"""Build the editable, evidence-grounded eight-page project case study.

Usage: python scripts/build_case_study.py [--pdf]
Requires: python-docx; optional PDF export uses an installed LibreOffice soffice. Read-only inputs: reports/, data/, source code.
Rendering: LibreOffice PDF export, followed by per-page visual review.
"""
from pathlib import Path
import json
import argparse
import shutil
import subprocess
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs'
OUT.mkdir(exist_ok=True)
METRICS = json.loads((ROOT / 'reports/metrics.json').read_text())
POLICY = json.loads((ROOT / 'reports/operating_policy.json').read_text())
WORKFLOW = json.loads((ROOT / 'reports/workflow_metrics.json').read_text())
SPLIT = json.loads((ROOT / 'reports/split_audit.json').read_text())
BASE = METRICS['models']['Word TF-IDF logistic regression']
CAND = METRICS['models']['Word + character calibrated SVM']
NAVY = '163047'
TEAL = '427C92'
AMBER = 'D98D32'
GRAY = '52616B'
PALE = 'F1F5F8'
BLACK = '000000'

D = Document()
sec = D.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(.65)
sec.bottom_margin = Inches(.64)
sec.left_margin = Inches(.76)
sec.right_margin = Inches(.76)
sec.header_distance = Inches(.27)
sec.footer_distance = Inches(.27)

styles = D.styles
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
    s = styles[name]
    s.font.name = 'Arial'
    s.font.color.rgb = RGBColor.from_string(BLACK)
    s.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.line_spacing = 1.13
styles['Normal'].paragraph_format.space_after = Pt(7)
styles['Title'].font.size = Pt(26)
styles['Title'].font.bold = True
styles['Title'].paragraph_format.space_after = Pt(12)
styles['Title'].paragraph_format.line_spacing = 1.03
styles['Subtitle'].font.size = Pt(12)
styles['Subtitle'].font.italic = False
styles['Subtitle'].paragraph_format.space_after = Pt(15)
styles['Heading 1'].font.size = Pt(20)
styles['Heading 1'].font.bold = True
styles['Heading 1'].paragraph_format.space_before = Pt(0)
styles['Heading 1'].paragraph_format.space_after = Pt(11)
styles['Heading 2'].font.size = Pt(12)
styles['Heading 2'].font.bold = True
styles['Heading 2'].paragraph_format.space_before = Pt(10)
styles['Heading 2'].paragraph_format.space_after = Pt(5)
styles['Caption'].font.size = Pt(9)
styles['Caption'].font.color.rgb = RGBColor.from_string(GRAY)
styles['Caption'].paragraph_format.space_after = Pt(9)
styles['Caption'].paragraph_format.line_spacing = 1.04
for s in styles:
    if s.type == 1 and s.name.startswith('Heading'):
        s.font.color.rgb = RGBColor.from_string(BLACK)

D.core_properties.title = 'AI Support Ticket Triage and Quality Review System'
D.core_properties.subject = 'An evaluated assistive banking support triage demonstration'
D.core_properties.author = 'Keisha Walker'
D.core_properties.keywords = 'BANKING77, intent classification, selective review, extractive summary'
D.core_properties.comments = ''

header = sec.header.paragraphs[0]
header.paragraph_format.space_after = Pt(0)
r=header.add_run('KEISHA WALKER  |  AI SUPPORT TICKET TRIAGE')
r.font.name='Arial'; r.font.size=Pt(8); r.font.color.rgb=RGBColor(0,0,0)
footer = sec.footer.paragraphs[0]
footer.paragraph_format.space_after=Pt(0)
r=footer.add_run('Case study  |  October 2026')
r.font.name='Arial'; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)
footer.paragraph_format.tab_stops.clear_all()
footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.95), WD_ALIGN_PARAGRAPH.RIGHT)
footer.add_run('\t')
r=footer.add_run(); r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.addnext(fld)


def para(text='', bold_prefix=None, style=None, after=None):
    p=D.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold=True; p.add_run(text[len(bold_prefix):])
    else: p.add_run(text)
    if after is not None: p.paragraph_format.space_after=Pt(after)
    p.paragraph_format.widow_control=True
    return p


def h(text):
    return D.add_paragraph(text, 'Heading 2')


def page(num, title, lead=None):
    p=para(f'{num:02d}  CASE STUDY', after=7)
    p.paragraph_format.page_break_before=True
    p.runs[0].font.size=Pt(8); p.runs[0].font.bold=True
    p.runs[0].font.color.rgb=RGBColor.from_string(GRAY)
    D.add_paragraph(title, 'Heading 1')
    if lead: para(lead)


def cell_fill(cell, color):
    pr=cell._tc.get_or_add_tcPr()
    shd=OxmlElement('w:shd'); shd.set(qn('w:fill'), color); pr.append(shd)


def table(headers, rows, widths, font=9.5):
    t=D.add_table(rows=1, cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side,n in [('top',90),('bottom',90),('left',110),('right',110)]:
        e=OxmlElement('w:'+side); e.set(qn('w:w'),str(n)); e.set(qn('w:type'),'dxa'); margins.append(e)
    pr.append(margins)
    for i,head in enumerate(headers): t.rows[0].cells[i].text=head
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=str(v)
    for j,row in enumerate(t.rows):
        cant=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
        for i,cell in enumerate(row.cells):
            cell.width=Inches(widths[i]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_fill(cell,NAVY if j==0 else (PALE if j%2==0 else 'FFFFFF'))
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
                p.paragraph_format.line_spacing=1.05
                if i>0 and len(str(cell.text))<22: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name='Arial'; r.font.size=Pt(font)
                    r.font.bold=(j==0); r.font.color.rgb=RGBColor.from_string('FFFFFF' if j==0 else BLACK)
    para('',after=1).paragraph_format.line_spacing=0.5
    return t


def figure(name, caption, width=6.8):
    p=D.add_paragraph(); p.paragraph_format.space_after=Pt(3); p.paragraph_format.keep_with_next=True
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(); r.add_picture(str(ROOT/'reports/figures'/name),width=Inches(width))
    inline=r._r.xpath('.//wp:docPr')
    if inline: inline[0].set('descr',caption)
    para(caption,style='Caption')


def small(text):
    p=para(text)
    for r in p.runs: r.font.size=Pt(9)
    p.paragraph_format.line_spacing=1.06
    return p


def command(text):
    p=para(text,after=4)
    p.paragraph_format.line_spacing=1.04
    for r in p.runs: r.font.name='DejaVu Sans Mono'; r.font.size=Pt(9)
    return p

# PAGE 1
p=para('APPLIED MACHINE LEARNING CASE STUDY',after=15)
p.runs[0].font.size=Pt(9); p.runs[0].font.bold=True
D.add_paragraph('AI Support Ticket Triage\nand Quality Review System', 'Title')
D.add_paragraph('Keisha Walker\nAn evaluated assistive workflow for banking support messages', 'Subtitle')
para('The prototype suggests a support queue and intent, extracts a short source-based summary, and surfaces reasons for human review. Its strongest result is a reproducible intent-classification comparison, paired with explicit evidence of where the surrounding workflow still fails.')
para('Scope matters. This is a limited demonstration using the public BANKING77 English banking-intent benchmark. It is not an IT-ticket dataset, its production origin is not verified here, and the app performs no account or financial actions.',bold_prefix='Scope matters.')
h('Results at a glance')
table(['Measure','Observed result','What it establishes'],[
    ['77-intent test macro-F1',f"{CAND['test']['macro_f1']:.3f}",'Higher than the 0.877 baseline on this fixed benchmark'],
    ['Full review policy',f"{POLICY['test_full_review_policy']['coverage']:.2%} coverage",'5.34% intent error among unflagged messages'],
    ['Summary phrase retention','82 of 99 phrases','Lexical diagnostic only; all phrases retained in 19 of 36 cases'],
    ['Software checks','26 tests passed','Checks implementation behavior, not real-world service quality'],
], [2.0,1.65,3.33],font=10)
h('Business problem and design response')
para('A support worker must understand the issue, choose an appropriate destination and preserve details for the next person. Misrouting, missing context and overconfident recommendations are plausible risks. The project makes those decisions inspectable without claiming that it has reduced handling time or improved customer outcomes.')
para('The interface keeps the original message beside the suggested intent, estimated confidence, alternatives, extractive summary and review reasons. A person owns every routing and resolution decision, including cases where no review rule fires.')
h('Main conclusion')
para('The richer word-and-character model is the better candidate in this experiment. The complete workflow remains a supervised demonstration: review heuristics reduce coverage without improving accepted-case error, and small diagnostic sets reveal missed off-domain and ambiguous requests. Production readiness would require independently reviewed target-domain data and operational evaluation.')
small('Evidence: frozen experiment outputs generated on 18 October 2026. References and reproduction instructions appear on page 8.')

# PAGE 2
page(2,'The assistive workflow','The system is intentionally narrow: a local, single-message suggestion tool rather than an autonomous agent or an integrated service desk.')
table(['Stage','Implementation','Human responsibility'],[
    ['1  Enter a message','Public or synthetic English banking-support text; 5,000-character cap','Remove private data before use'],
    ['2  Estimate the intent','77-class model; top three alternatives; probability for the top intent','Check what the message is actually asking'],
    ['3  Suggest a queue','Fixed mapping from 77 intents to nine project-defined queues','Confirm the queue and any specialist need'],
    ['4  Extract a summary','Rank and copy up to two sentences; cap at 60 words','Compare with the full source and missing context'],
    ['5  Inspect review reasons','Confidence, margin, domain, context and topic rules','Make the final routing and response decision'],
], [1.35,2.78,2.85],font=9.6)
h('What the interface makes visible')
para('The Streamlit app displays the suggested category, fine-grained intent, confidence, top alternatives and a review flag. Confidence means the model\'s estimated probability for the 77-class intent; it is not the probability that the broader queue is correct and is not a correctness guarantee. Users can download the result as JSON.')
para('A review flag is an explanation prompt, not a separate learned quality score. The checklist asks whether the queue fits the issue, whether the summary preserves negation and timing, whether multiple issues or a security concern appear, and whether a specialist should inspect the original.')
h('A transparent operational taxonomy')
para('The nine queues are Card lifecycle; Card payments; Top ups and deposits; Transfers; Cash and ATMs; Account access and profile; Identity verification; Foreign exchange and coverage; and Security and disputes. This mapping is defined in project code before evaluation. BANKING77 does not provide queue labels, routing outcomes or an organization\'s ownership model.')
h('Privacy and action boundaries')
para('The model runs in the local Python process. The application contains no model API calls, account access, automatic ticket routing or ticket persistence. Hosting the app elsewhere changes who can receive input, so the local-only assumption must be revisited before deployment.')
para('Basic masking catches email-like strings and long number sequences. It is not comprehensive de-identification or a data-loss-prevention control. Do not enter passwords, card numbers, identity documents or private customer records. The tool must not be used for credit, fraud adjudication, eligibility or account changes.')
para('Why this matters. An inspectable suggestion can support an operator, but a confident prediction or an unflagged message is never permission to act.',bold_prefix='Why this matters.')

# PAGE 3
page(3,'Benchmark and evaluation design','The experiment separates model fitting, model selection and final reporting, while documenting the limits of the benchmark itself.')
h('Data provenance and permitted claims')
para('BANKING77 by PolyAI contains 10,003 official training examples, 3,080 official test examples and 77 banking intents. The official test has 40 rows per intent. The source describes online-banking queries; the collection method and whether all queries originated in real customer interactions are not established here. It should not be called verified production ticket data or definitively synthetic data. [1-3]')
para('There are no source labels for urgency, priority, service-level agreements, handling time, resolution quality or customer satisfaction. The nine queues and review rules are project additions, not supervised ground truth for those business outcomes.')
h('Exact duplicate hygiene')
table(['Development data step','Rows removed or retained'],[
    ['Original official development split','10,003 rows'],
    ['Remove normalized text overlapping official test','25 removed'],
    ['Remove conflicting-label normalized strings','2 removed'],
    ['Remove remaining normalized development duplicates','30 removed'],
    ['Stratified split of remaining development data','7,956 train / 1,990 validation'],
    ['Official test rows, order and labels','3,080 retained unchanged'],
], [4.75,2.23],font=9.8)
para('Normalization lowercases text, replaces non-word runs with a space and strips edges. Official test text is used only for this disclosed exact-overlap hygiene, which removes development rows rather than altering the test set. No validation or test text fits a vocabulary or classifier.')
para('Residual similarity remains. Four normalized duplicate test rows are retained. A post-fit audit finds 40 test rows with word-TF-IDF cosine similarity of at least 0.95 to a training row. These observations limit independence claims; similar paraphrases can survive exact deduplication.',bold_prefix='Residual similarity remains.')
h('Selection and reporting discipline')
para('The development split is stratified by intent with seed 42. The model is selected on validation intent macro-F1; the review threshold is selected on validation coverage subject to an empirical error target. The official test is then used for final reporting. No model, threshold, summary rule or review rule is retuned after seeing test scores or diagnostic outcomes.')
small('Source records: data/DATA_PROVENANCE.md, data/raw/SOURCE_MANIFEST.json and reports/split_audit.json. The preliminary provenance audit used a simpler normalization; the table above reports the final implemented policy.')

# PAGE 4
page(4,'Model comparison and calibration','The comparison tests a simple, interpretable baseline against a richer sparse-text representation. It does not establish that one model family is inherently superior.')
figure('model_comparison.png','Figure 1. Model selection uses validation macro-F1; test results are reported after selection.',width=5.25)
table(['Metric','Word logistic baseline','Word and character SVM'],[
    ['Validation intent macro-F1',f"{BASE['validation']['macro_f1']:.6f}",f"{CAND['validation']['macro_f1']:.6f}"],
    ['Test intent macro-F1',f"{BASE['test']['macro_f1']:.6f}",f"{CAND['test']['macro_f1']:.6f}"],
    ['Test nine-queue macro-F1',f"{BASE['test']['queue_macro_f1']:.6f}",f"{CAND['test']['queue_macro_f1']:.6f}"],
    ['Test log loss',f"{BASE['test']['log_loss']:.3f}",f"{CAND['test']['log_loss']:.3f}"],
    ['Test calibration error in 10 bins',f"{BASE['test']['expected_calibration_error_10_bins']:.3f}",f"{CAND['test']['expected_calibration_error_10_bins']:.3f}"],
], [3.0,1.99,1.99],font=9.5)
h('What changed between models')
para('The baseline uses word TF-IDF unigrams and bigrams with logistic regression (C=4). The candidate combines the same word features with character-within-word 3-5-grams and a linear SVM (C=1). Sigmoid calibration uses three stratified training-only folds; the full vectorizer-and-classifier pipeline is cross-fitted to avoid leaking calibration-fold vocabulary.')
para('Character fragments provide a plausible way to represent spelling variation, but the experiment changes features, estimator and calibration together. It cannot isolate which change causes the observed gain. The 0.0267 test macro-F1 increase is evidence for this candidate configuration on this benchmark.')
para('The 0.956 queue macro-F1 is a separate, easier aggregation: several intent errors remain inside the same broad queue. It must not replace the 77-intent score. Lower calibration error and log loss are encouraging within-domain, but do not validate confidence on unseen domains.')
small('Uncertainty: a 500-replicate paired row bootstrap gives a 95% percentile interval of 0.892-0.912 for candidate test macro-F1 and 0.0185-0.0353 for its improvement. These intervals omit training variability, domain shift and correlated paraphrases.')

# PAGE 5
page(5,'Selective review and actual coverage','The app combines a confidence threshold with additional fixed review rules. The full policy must be evaluated separately from the confidence-only curve.')
figure('coverage_error.png','Figure 2. Each curve varies only the confidence threshold. It does not represent the full app policy.',width=6.1)
table(['Policy and split','Accepted','Coverage','Intent error among accepted'],[
    ['Confidence only on validation','1,797 / 1,990','90.30%','88 / 1,797 = 4.90%'],
    ['Confidence only on test','2,788 / 3,080','90.52%','140 / 2,788 = 5.02%'],
    ['Full app policy on test','2,152 / 3,080','69.87%','115 / 2,152 = 5.34%'],
], [2.5,1.3,1.0,2.18],font=9.5)
h('How the operating point is selected')
para('Threshold 0.57 maximizes validation coverage among 0.01-grid thresholds with empirical accepted-case error at or below 5% and at least 100 accepted examples. Ties use the lower threshold; if none qualifies, the fallback is review all. The 5% target is a selection criterion, not a guaranteed test or production error bound.')
para('The full policy also flags a top-two intent margin below 0.10, weak banking cues, fewer than two known terms, fewer than four words, possible multiple issues, security or dispute topics, and summary truncation. “Accepted” means no rule fired, not that the system autonomously acted.')
para('Important negative result. The full policy reviews 928 of 3,080 test messages and leaves 115 errors among 2,152 unflagged messages. Its unflagged subset contains fewer total errors, but the accepted-case error rate is higher than confidence-only filtering. These heuristics do not demonstrate improved selective error on this test set.',bold_prefix='Important negative result.')
small('95% Wilson intervals for test accepted-case error: confidence-only 4.27%-5.90%; full policy 4.47%-6.38%. Neither interval is a guarantee under domain shift. The operating policy remains unchanged after this evaluation.')

# PAGE 6
page(6,'Summary and review diagnostics','Project-authored synthetic fixtures expose failure modes without pretending to be human-rated, externally validated or representative customer data.')
h('Extractive summary method')
para('The summarizer ranks sentences using issue cues, banking-domain cues and a bounded length term, then copies at most two sentences in source order. It caps output at 60 whitespace-separated words and marks truncation. One-sentence inputs may be unchanged. Sentence splitting is rule-based and can mishandle abbreviations.')
para('Thirty-six independently authored fixture messages contain 99 expected key phrases. Each phrase is a case-insensitive contiguous source span. Expectations were frozen before inference and were not supplied by human raters. The comparator simply copies the first two sentences.')
table(['Summary diagnostic','Lead two sentences','Issue-focused extractor'],[
    ['Expected phrases retained by micro average','51 / 99 = 51.52%','82 / 99 = 82.83%'],
    ['All expected phrases retained','7 / 36 = 19.44%','19 / 36 = 52.78%'],
    ['Summaries fully copied from source','36 / 36','36 / 36'],
    ['Mean summary / input word count','51.73%','57.58%'],
], [3.15,1.91,1.92],font=9.5)
para('The issue-focused method retains more authored phrases, partly at greater output length. It still misses at least one expected phrase in 17 of 36 messages. Exact phrase retention is not semantic correctness, factual accuracy or actionability. Fully extractive output is primarily a construction check: copying can preserve a false user statement or omit a critical negation or qualifying sentence.')
h('Review challenge results')
table(['Authored challenge group','Flagged','Missed or mismatched'],[
    ['Off-domain requests','7 / 8','1 not flagged'],
    ['Ambiguous or multiple-issue requests','5 / 8','3 not flagged'],
    ['Straightforward in-domain requests','1 / 8','1 flagged against expected no-review'],
], [3.55,1.0,2.43],font=9.5)
para('The ambiguous group mixes four insufficient-context cases and four clearly multi-intent cases. The 5-of-8 result is for that combined group, not a pure multi-issue recall estimate. The one straightforward mismatch is a lost-card request: the authored expectation was no review, while the fixed security rule escalates it. That mismatch reflects a policy disagreement as well as a diagnostic result.')
h('How to interpret this evidence')
para('These fixtures were authored by an AI assistant for the project, not sampled from customer records. They are small, intentionally selected tests with lexical traps and difficult issue placement. No model, threshold, extractor or review rule was tuned after their scores were inspected. A new independent set is needed if development responds to these failures.')
small('Detailed outputs and missing spans are retained in reports/summary_diagnostics.csv and reports/review_diagnostics.csv; authoring methods and limitations are in data/FIXTURES.md.')

# PAGE 7
page(7,'Failure analysis and next validation','The most useful next step is independent validation on an authorized target workflow, not a broader claim based on the current benchmark.')
h('A confident off-domain failure')
para('The synthetic input “My graphics card is not working after a computer restart. The screen stays black.” is assigned card_not_working with 86.78% intent confidence. No review rule fires. The word “card” satisfies the banking cue check, and the closed-set classifier cannot represent a separate out-of-domain class.')
para('This is an observed exploratory failure, not robust out-of-domain detection. Similar lexical traps can look familiar to the model while referring to a different task. The example remains in the diagnostic record; rules were not patched to make the demonstration look stronger.')
h('Confident errors inside the benchmark')
table(['Official test example','Expected and predicted intent','Lesson'],[
    ['“can I have a non virtual card?”','order_physical_card → getting_virtual_card\n97.26% confidence; no flag','Negation can reverse the intended request'],
    ['“My transfer is pending.”','balance_not_updated_after_bank_transfer → pending_transfer\n93.55% confidence; no flag','Short text can cross queue boundaries'],
    ['“Where do I find the exchange rate?”','exchange_charge → exchange_rate\n95.12% confidence; no flag','Benchmark labels can also be hard to infer from surface text'],
], [2.03,3.0,1.95],font=9.3)
para('These are disagreements with benchmark labels, not adjudicated service errors. The exchange example especially warrants label review before treating it as model misunderstanding. Other frequent confusions involve card disputes, failed versus reverted top-ups and transfer status. The full error log keeps source text, both labels, confidence and review reasons.')
h('A staged validation plan')
para('1. Establish the intended domain. Obtain permission to use de-identified target-domain messages. Define queue ownership, multi-issue handling and meaningful error costs with support staff. Create new labels through independent review and adjudication.',bold_prefix='1. Establish the intended domain.')
para('2. Evaluate on a fresh holdout. Use time- or source-separated data with stronger paraphrase audits. Measure per-intent and per-queue performance, rejection behavior, false escalations, missed security concerns, multi-intent cases and off-domain traps.',bold_prefix='2. Evaluate on a fresh holdout.')
para('3. Test usefulness with people. Have reviewers assess summary completeness, negation, timing and actionability. Compare against simpler interfaces in a controlled study before estimating time savings or quality improvements.',bold_prefix='3. Test usefulness with people.')
para('4. Add deployment controls. Require privacy and security review, access controls, retention rules, auditability, monitoring and human escalation before any live integration. Keep account and financial actions outside this prototype.',bold_prefix='4. Add deployment controls.')

# PAGE 8
page(8,'Reproduction and source references','The project includes executable source, a notebook, a local demo, a saved model and row-level evaluation artifacts. No paid model API or external model credential is required.')
h('Run the local demonstration')
para('Use Python 3.12 in a fresh virtual environment. From the project folder, install the pinned runtime dependencies and launch Streamlit. The bundled model is compatible with the recorded dependency versions; retrain if loading fails.')
command('python -m venv .venv')
command('source .venv/bin/activate    # Windows: .venv\\Scripts\\activate')
command('python -m pip install -r requirements.txt')
command('streamlit run app.py')
para('Use only public or synthetic messages. Load only a trusted joblib model: deserialization can execute code. Do not expose the demo publicly without a deployment and privacy review.')
h('Reproduce the evidence')
command('python -m pip install -r requirements-dev.txt')
command('python -m src.train')
command('python -m src.evaluate_workflow')
command('python -m pytest -q')
para('The recorded run uses Python 3.12.14, scikit-learn 1.8.0 and seed 42. All 26 software tests pass. The executed notebook, per-class results, threshold curves, split audit, bootstrap intervals and diagnostic CSVs support review of the results. Exact runtimes can vary by machine; software tests do not establish production performance.')
h('Evidence map')
small('Classifier and calibration results: reports/metrics.json. Threshold and full-policy behavior: reports/operating_policy.json. Fixture diagnostics: reports/workflow_metrics.json. Dataset hygiene: reports/split_audit.json. Row-level failures: reports/error_analysis.csv. Reproducible analysis: notebooks/01_support_ticket_triage.ipynb.')
h('Sources and licensing')
small('[1] PolyAI. BANKING77, task-specific-datasets repository. Pinned commit 57ec275d8078af65b7731c2a98be812d844a6d6b; retrieved 18 October 2026. https://github.com/PolyAI-LDN/task-specific-datasets')
small('[2] Casanueva, I., Temcinas, T., Gerz, D., Henderson, M. and Vulić, I. (2020). Efficient Intent Detection with Dual Sentence Encoders. 2nd Workshop on NLP for ConvAI. https://arxiv.org/abs/2003.04807')
small('[3] PolyAI. BANKING77 dataset card. https://huggingface.co/datasets/PolyAI/banking77')
small('BANKING77 and its accompanying source materials retain the Creative Commons Attribution 4.0 International license: https://creativecommons.org/licenses/by/4.0/. Raw files are unmodified. Processed splits, the nine-queue mapping, review rules and synthetic fixtures are project additions. Attribution and checksums are preserved with the project.')
small('Project-authored code, documentation and synthetic fixtures are MIT licensed, copyright 2026 Keisha Walker. Dataset and dependency licenses remain separate. These results describe benchmark performance and development work, not deployed business impact or independent external validation.')

# Make external references clickable without changing the visible citation text.
for p in D.paragraphs:
    if 'https://' not in p.text:
        continue
    for run in list(p.runs):
        source=run.text
        matches=list(re.finditer(r'https://[^\s]+', source))
        if not matches:
            continue
        run.text=''
        cursor=0
        for match in matches:
            if match.start()>cursor:
                r=p.add_run(source[cursor:match.start()]); r.font.name='Arial'; r.font.size=Pt(9)
            url=match.group(0).rstrip('.;,')
            link=OxmlElement('w:hyperlink')
            rel=p.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
            link.set(qn('r:id'),rel)
            r=OxmlElement('w:r'); props=OxmlElement('w:rPr')
            color=OxmlElement('w:color'); color.set(qn('w:val'),TEAL); props.append(color)
            size=OxmlElement('w:sz'); size.set(qn('w:val'),'18'); props.append(size)
            r.append(props); text=OxmlElement('w:t'); text.text=url; r.append(text); link.append(r); p._p.append(link)
            suffix=match.group(0)[len(url):]
            if suffix:
                r=p.add_run(suffix); r.font.name='Arial'; r.font.size=Pt(9)
            cursor=match.end()
        if cursor<len(source):
            r=p.add_run(source[cursor:]); r.font.name='Arial'; r.font.size=Pt(9)

# Avoid accidental inherited title borders and theme colors.
for st in D.styles:
    for node in list(st.element.iter(qn('w:pBdr'))):
        node.getparent().remove(node)
for p in D.paragraphs:
    if p.style.name in {'Title','Subtitle','Heading 1','Heading 2'}:
        pr=p._p.get_or_add_pPr()
        for node in list(pr):
            if node.tag==qn('w:pBdr'): pr.remove(node)
        for r in p.runs: r.font.color.rgb=RGBColor.from_string(BLACK)
D.save(OUT/'Keisha_Walker_AI_Support_Ticket_Triage_Case_Study.docx')
print(OUT/'Keisha_Walker_AI_Support_Ticket_Triage_Case_Study.docx')

if __name__ == '__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--pdf',action='store_true',help='Also export PDF with installed LibreOffice')
    args=parser.parse_args()
    if args.pdf:
        soffice=shutil.which('soffice')
        if not soffice:
            raise SystemExit('DOCX saved. Install LibreOffice or export the DOCX as PDF from Word.')
        subprocess.run([soffice,'--headless','--convert-to','pdf','--outdir',str(OUT),str(OUT/'Keisha_Walker_AI_Support_Ticket_Triage_Case_Study.docx')],check=True)
